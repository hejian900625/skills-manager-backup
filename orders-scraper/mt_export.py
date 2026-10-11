# 美团订单导出（统一订单中心 H5：i.meituan.com/mttouch/page/order）
#
# 原理（与京东/拼多多同一套路：真实浏览器操作 + 监听接口）：
#   ① 首屏：页面 SSR 内嵌 `window.__INITIAL_STATE__.orderList`（约 20 条，全字段结构化）；
#   ② 翻页：下拉触发 `GET /mttouch/api/order/getOrderList?lastOrderTime=<游标>&status=0...`，
#      游标是 lastOrderTime（首页数据里的顶层字段）。
#   ③ 登录态存 `数据结果/mt_session.json`（手机号+验证码，约数天有效）。
#
# 已知限制（实测结论）：
#   · 订单金额没有结构化字段，只有 orderinfo 文本（「实付: ￥86」「总价：¥33.16」）→ 正则解析。
#   · 美团有 H5guard + dfp 设备指纹风控，翻页节奏太快可能被限流 → 每轮之间多等一会。
#
# 跑法（在 C:/Users/Lenovo 下跑，产出落到 数据结果/）：
#   C:/Users/Lenovo/.workbuddy/binaries/python/versions/3.13.12/python.exe "C:/Users/Lenovo/.workbuddy/skills/orders_scraper/mt_export.py"
import csv
import datetime as dt
import json
import os
import re
import sys
import time
import traceback

from playwright.sync_api import sync_playwright

OUT = "数据结果"
API_DIR = os.path.join(OUT, "mt_api")
ORDER_URL = "https://i.meituan.com/mttouch/page/order"
STATE = os.path.join(OUT, "mt_session.json")
LOG_PATH = os.path.join(OUT, "mt_export.log")

H5_UA = ("Mozilla/5.0 (iPhone; CPU iPhone OS 17_0 like Mac OS X) AppleWebKit/605.1.15 "
         "(KHTML, like Gecko) Version/17.0 Mobile/15E148 Safari/604.1")

LOGIN_COOKIE_HINTS = ("token", "userid", "user_id", "mt_c_token")

ORDERS = {}
REQ_COUNT = {"n": 0}
RAW_IDX = {"n": 0}
PAGE_SIZES = []
BAD_ROWS = {"n": 0}
ERR_COUNT = {"n": 0}
LAST_CURSOR = {"v": None}
# 接口是否已经返回空页（=真到底）
EMPTY_PAGE = {"hit": False}


class _Tee:
    def __init__(self, *streams):
        self.streams = streams

    def write(self, s):
        for f in self.streams:
            try:
                f.write(s)
            except Exception:
                pass

    def flush(self):
        for f in self.streams:
            try:
                f.flush()
            except Exception:
                pass


def pump(page, ms):
    """等待并驱动 Playwright 事件循环（裸 time.sleep 不驱动，响应回调不会执行）。"""
    try:
        page.wait_for_timeout(ms)
    except Exception:
        time.sleep(ms / 1000.0)


def fmt_time(ts):
    try:
        ts = int(ts)
        if ts <= 0:
            return ""
        return dt.datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return ""


# ---------------------------------------------------------------- 字段抽取
# orderinfo 里这几类是「元信息」而不是商品描述（外卖订单的 orderinfo 常常只有这两行）
META_LINE = re.compile(r"^(下单时间|实付|总价|金额|付款|共[省减]|优惠|配送|预计)")

FIELDS = ["订单号", "下单时间", "订单状态", "商家", "商品", "实付金额", "金额原文",
          "订单分类", "partnerid", "详情链接", "来源"]


def parse_amount(info_lines):
    """从 orderinfo 文本里抠金额：优先「实付」，其次「总价」「金额」「售价」。"""
    pats = [r"实付[:：]\s*[¥￥]\s*([\d,]+(?:\.\d+)?)",
            r"总价[:：]\s*[¥￥]\s*([\d,]+(?:\.\d+)?)",
            r"金额[:：]\s*[¥￥]\s*([\d,]+(?:\.\d+)?)",
            r"[¥￥]\s*([\d,]+(?:\.\d+)?)"]
    for p in pats:
        for line in info_lines:
            m = re.search(p, line)
            if m:
                v = m.group(1).replace(",", "")
                try:
                    return f"{float(v):.2f}", line.strip()
                except Exception:
                    return v, line.strip()
    return "", ""


def parse_order_time(o, info_lines):
    t = o.get("ordertime")
    if t:
        s = fmt_time(t)
        if s:
            return s
    for line in info_lines:
        m = re.search(r"下单时间[:：]\s*([0-9]{4}-[0-9]{2}-[0-9]{2}[ T]?[0-9:]*)", line)
        if m:
            return m.group(1).strip()
    return ""


def walk_orders(obj, sink):
    """递归找出所有「像美团订单」的 dict（含 orderid / stringOrderId）。"""
    if isinstance(obj, dict):
        if obj.get("orderid") or obj.get("stringOrderId") or obj.get("orderId"):
            sink.append(obj)
        for v in obj.values():
            walk_orders(v, sink)
    elif isinstance(obj, list):
        for v in obj:
            walk_orders(v, sink)


def iter_order_objs(obj):
    """优先取顶层 orderList / data.orderList，再 walk 兜底。"""
    if isinstance(obj, dict):
        ol = obj.get("orderList")
        if isinstance(ol, list) and ol:
            return ol
        data = obj.get("data")
        if isinstance(data, dict) and isinstance(data.get("orderList"), list):
            return data["orderList"]
        if isinstance(data, list) and data:
            return data
    sink = []
    walk_orders(obj, sink)
    return sink


def order_to_row(o):
    if not isinstance(o, dict):
        return None
    try:
        raw_id = str(o.get("orderid") or o.get("orderId") or "")
        sn = str(o.get("stringOrderId") or o.get("orderViewId") or raw_id)
        if not sn:
            return None
        info = o.get("orderinfo")
        if not isinstance(info, list):
            info = []
        info = [str(x) for x in info]

        amount, amount_src = parse_amount(info)
        # 商品描述：跳过「下单时间 / 实付 / 总价」这类元信息行
        goods_lines = [x.strip() for x in info if x.strip() and not META_LINE.search(x.strip())]
        goods = " / ".join(goods_lines)
        cate = o.get("catename") or ""
        pid = o.get("partnerid")

        row = {
            "订单号": sn,
            "下单时间": parse_order_time(o, info),
            "订单状态": o.get("showstatus") or "",
            "商家": o.get("title") or "",
            "商品": goods,
            "实付金额": amount,
            "金额原文": amount_src,
            "订单分类": cate,
            "partnerid": "" if pid is None else str(pid),
            "详情链接": o.get("orderDetail") or "",
            "来源": "美团H5",
        }
        return row
    except Exception as e:
        BAD_ROWS["n"] += 1
        if BAD_ROWS["n"] <= 3:
            print(f"   !! 单条订单解析失败（已跳过）: {type(e).__name__}: {e}")
        return None


def collect(obj, tag=""):
    sink = iter_order_objs(obj)
    PAGE_SIZES.append(len(sink))
    added = 0
    for o in sink:
        row = order_to_row(o)
        if not row:
            continue
        sn = row["订单号"]
        if sn in ORDERS:
            for k, v in row.items():
                if v not in ("", None) and not ORDERS[sn].get(k):
                    ORDERS[sn][k] = v
            continue
        ORDERS[sn] = row
        added += 1
    return added


def read_initial_state(page):
    """读 SSR 内嵌的 window.__INITIAL_STATE__（含首屏 orderList + lastOrderTime 游标）。"""
    try:
        st = page.evaluate("() => window.__INITIAL_STATE__ || null")
        if isinstance(st, dict):
            cur = st.get("lastOrderTime")
            if cur:
                LAST_CURSOR["v"] = cur
            return st
    except Exception:
        pass
    return None


def has_login_cookie(ctx):
    try:
        names = {c.get("name", "").lower() for c in ctx.cookies()}
        return any(any(h in n for h in LOGIN_COOKIE_HINTS) for n in names)
    except Exception:
        return False


def do_scroll(page):
    """滚到底。美团订单页是超长 H5 列表（900 条约 9 万像素高），
    每轮必须滚得够多、且给足 lazy-load 时间，否则会在中途「假到底」。"""
    try:
        box = page.viewport_size or {"width": 390, "height": 844}
        page.mouse.move(box["width"] / 2, box["height"] / 2)
        for _ in range(4):
            page.mouse.wheel(0, 3000)
            pump(page, 300)
    except Exception:
        pass
    try:
        page.evaluate("""() => {
            const els = [document.scrollingElement, document.body, document.documentElement];
            document.querySelectorAll('*').forEach(e => {
                if (e.scrollHeight > e.clientHeight + 200) els.push(e);
            });
            els.forEach(e => { try { e.scrollTop = e.scrollHeight; } catch(_){} });
            window.scrollTo(0, document.body.scrollHeight);
            window.dispatchEvent(new Event('scroll'));
        }""")
    except Exception:
        pass


def no_more_visible(page):
    try:
        return bool(page.evaluate("""() => {
            const t = document.body ? document.body.innerText : '';
            return /没有更多|已经到底|无更多订单/.test(t);
        }"""))
    except Exception:
        return False


def fetch_page(page, cursor):
    """在页面上下文里请求「下一页」——同源 fetch 会自动带 cookie。

    接口是 lastOrderTime 游标分页，顺着游标走比滚动可靠得多。
    """
    try:
        return page.evaluate("""async (cursor) => {
            const url = '/mttouch/api/order/getOrderList?lat=0&lng=0&lastOrderTime='
                + cursor + '&status=0&shark=1&yodaReady=h5&csecplatform=4&csecversion=4.3.0';
            try {
                const r = await fetch(url, {credentials: 'include',
                                            headers: {'accept': 'application/json'}});
                return {status: r.status, text: await r.text()};
            } catch (e) {
                return {status: -1, text: String(e)};
            }
        }""", cursor)
    except Exception as e:
        return {"status": -2, "text": f"{type(e).__name__}: {e}"}


# partnerid → 业务线。优先用「已有分类的订单」自动实证，学不到的用下表兜底。
# 下表依据订单内容特征判定（2026-10-10 实测样本）：
#   137=「电单车/开锁时间」共享单车 · 68=密室/KTV 预订 · 136=「已过号」餐厅排队取号
#   54/151=券包 · 51=演出票务（演唱会）
KNOWN_PID = {
    "54": "券包", "151": "券包", "51": "演出票务",
    "68": "到店预订", "136": "排队取号", "137": "共享单车",
}


def fill_categories():
    """补齐缺失的「订单分类」：① 用已有分类的订单实证 partnerid→分类 映射；② 再用 KNOWN_PID 兜底。"""
    learned = {}
    for r in ORDERS.values():
        pid, cate = r.get("partnerid"), r.get("订单分类")
        if pid and cate:
            learned.setdefault(pid, {})
            learned[pid][cate] = learned[pid].get(cate, 0) + 1
    best = {k: max(v.items(), key=lambda x: x[1])[0] for k, v in learned.items()}
    filled = 0
    for r in ORDERS.values():
        if r.get("订单分类"):
            continue
        cate = best.get(r.get("partnerid") or "") or KNOWN_PID.get(r.get("partnerid") or "")
        if cate:
            r["订单分类"] = cate
            filled += 1
    return filled


# ---------------------------------------------------------------- 导出
def write_files():
    rows = sorted(ORDERS.values(), key=lambda r: (r.get("下单时间") or ""), reverse=True)
    date = dt.datetime.now().strftime("%Y-%m-%d")
    base = os.path.join(OUT, f"mt_orders_{date}")
    csv_path = base + ".csv"
    xlsx_path = base + ".xlsx"

    def clean(v):
        return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()

    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(FIELDS)
        for r in rows:
            w.writerow([clean(r.get(k, "")) for k in FIELDS])
    print(f">>> CSV 已写出: {csv_path}（{len(rows)} 行）")

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "美团订单"
        ws.append(FIELDS)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="FFC300")   # 美团黄
            c.alignment = Alignment(horizontal="center")
        for r in rows:
            ws.append([clean(r.get(k, "")) for k in FIELDS])
        widths = [26, 20, 12, 26, 44, 12, 22, 14, 12, 56, 10]
        for i, wd in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = wd
        ws.freeze_panes = "A2"
        wb.save(xlsx_path)
        print(f">>> Excel 已写出: {xlsx_path}")
    except Exception as e:
        print("!! 写 Excel 失败:", type(e).__name__, e)
    return rows


def summary(rows):
    print()
    print("=" * 62)
    print(f"共解析订单: {len(rows)} 条")
    if not rows:
        return
    years, status, cate = {}, {}, {}
    total, noamount = 0.0, 0
    for r in rows:
        y = (r.get("下单时间") or "")[:4] or "(无时间)"
        years[y] = years.get(y, 0) + 1
        s = r.get("订单状态") or "(无状态)"
        status[s] = status.get(s, 0) + 1
        c = r.get("订单分类") or "(无分类)"
        cate[c] = cate.get(c, 0) + 1
        if r.get("实付金额"):
            try:
                total += float(r["实付金额"])
            except Exception:
                pass
        else:
            noamount += 1
    print("年份分布:", "  ".join(f"{k}:{v}" for k, v in sorted(years.items(), reverse=True)))
    print("状态分布:", "  ".join(f"{k}:{v}" for k, v in status.items()))
    print("分类分布:", "  ".join(f"{k}:{v}" for k, v in
                              sorted(cate.items(), key=lambda x: -x[1])[:12]))
    print(f"实付金额合计: ¥{total:.2f}（{noamount} 条无金额）")
    ts = [r["下单时间"] for r in rows if r.get("下单时间")]
    if ts:
        print("时间跨度:", min(ts)[:10], "~", max(ts)[:10])
    print("=" * 62)


# ---------------------------------------------------------------- 主流程
def main():
    os.makedirs(OUT, exist_ok=True)
    os.makedirs(API_DIR, exist_ok=True)
    have = os.path.exists(STATE)
    print(f">>> 登录态文件: {STATE} （{'已存在，将复用' if have else '不存在，需登录'}）")

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False,
                                    args=["--disable-blink-features=AutomationControlled"])
        kw = {"user_agent": H5_UA, "viewport": {"width": 390, "height": 844},
              "is_mobile": True, "has_touch": True,
              "locale": "zh-CN", "timezone_id": "Asia/Shanghai"}
        ctx = browser.new_context(storage_state=STATE, **kw) if have else browser.new_context(**kw)
        page = ctx.new_page()
        page.set_default_timeout(60000)

        def on_response(resp):
            try:
                u = resp.url
                if "getOrderList" not in u:
                    return
                data = None
                try:
                    data = resp.json()
                except Exception:
                    try:
                        txt = resp.text()
                        m = re.search(r"\{.*\}", txt, re.S)
                        if m:
                            data = json.loads(m.group(0))
                    except Exception:
                        data = None
                if data is None:
                    return
                REQ_COUNT["n"] += 1
                RAW_IDX["n"] += 1
                try:
                    with open(os.path.join(API_DIR, f"raw_{RAW_IDX['n']:02d}.json"),
                              "w", encoding="utf-8") as f:
                        json.dump(data, f, ensure_ascii=False, indent=1)
                except Exception:
                    pass
                n = collect(data)
                page_n = PAGE_SIZES[-1] if PAGE_SIZES else 0
                if page_n == 0:
                    EMPTY_PAGE["hit"] = True
                print(f"[API] getOrderList 第 {REQ_COUNT['n']} 次：本页 {page_n} 条，"
                      f"新增 {n}，累计 {len(ORDERS)}"
                      + ("  ← 空页，已到底" if page_n == 0 else ""))
            except Exception as e:
                ERR_COUNT["n"] += 1
                if ERR_COUNT["n"] <= 5:
                    print(f"!! 响应处理异常: {type(e).__name__}: {e}")

        page.on("response", on_response)

        for attempt in range(3):
            try:
                page.goto(ORDER_URL, wait_until="domcontentloaded")
                break
            except Exception as e:
                print(f"!! goto 第 {attempt+1} 次失败: {type(e).__name__}")
                pump(page, 2000)
        pump(page, 5000)
        print(">>> 当前 URL:", page.url)

        if not has_login_cookie(ctx) or "passport.meituan.com" in page.url:
            print(">>> 请在窗口里用【手机号 + 短信验证码】登录（最多等 300 秒）…")
            deadline = time.time() + 300
            last = -1
            while time.time() < deadline:
                if has_login_cookie(ctx) and "passport.meituan.com" not in page.url:
                    print(">>> 检测到登录成功")
                    break
                w = int(300 - (deadline - time.time()))
                if w // 10 != last:
                    last = w // 10
                    print(f"    [{w}s] 等待登录… URL={page.url[:90]}")
                pump(page, 2000)
            else:
                print("!! 超时未登录，退出")
                try:
                    browser.close()
                except Exception:
                    pass
                return
            try:
                ctx.storage_state(path=STATE)
                print(">>> 登录态已保存:", STATE)
            except Exception:
                pass
            for _ in range(3):
                try:
                    page.goto(ORDER_URL, wait_until="domcontentloaded")
                    break
                except Exception:
                    pump(page, 2000)
            pump(page, 6000)
        else:
            print(">>> 登录态有效，免登录")
            try:
                ctx.storage_state(path=STATE)
            except Exception:
                pass

        print(">>> 已进入订单页:", page.url)
        try:
            print(">>> 页面标题:", page.title())
        except Exception:
            pass

        # 首屏（SSR 内嵌）
        st = read_initial_state(page)
        if st:
            n = collect(st, "initial")
            print(f">>> SSR __INITIAL_STATE__：新增 {n}，累计 {len(ORDERS)}"
                  f"（游标 lastOrderTime={LAST_CURSOR['v']}）")
        else:
            print("!! 未能读到 __INITIAL_STATE__")
        print(f">>> 首屏累计订单: {len(ORDERS)}")

        # ---- 翻页主路径：接口游标翻页 ----
        # 实测滚动方案在美团这里不稳定（滚太快页面不触发加载 → 提前「假到底」；
        # 滚太慢又慢），而接口本身就是 lastOrderTime 游标分页 → 顺游标请求最可靠。
        try:
            page.remove_listener("response", on_response)   # 避免同一响应被计两次
        except Exception:
            pass

        cursor = LAST_CURSOR["v"]
        print(f">>> 开始接口翻页，初始游标 lastOrderTime={cursor}")
        rnd, MAX_ROUNDS = 0, 300
        api_fail = 0
        while rnd < MAX_ROUNDS and cursor:
            rnd += 1
            res = fetch_page(page, cursor)
            txt = (res or {}).get("text") or ""
            data = None
            try:
                data = json.loads(txt)
            except Exception:
                data = None
            if not isinstance(data, dict) or data.get("code") != 0:
                api_fail += 1
                code = data.get("code") if isinstance(data, dict) else None
                print(f"   [{rnd}] 接口异常 HTTP {(res or {}).get('status')} code {code}，"
                      f"前 120 字 {txt[:120]}")
                if api_fail >= 3:
                    print("   连续 3 次接口异常 → 改用滚动兜底")
                    break
                pump(page, 3000)
                continue
            api_fail = 0
            REQ_COUNT["n"] += 1
            RAW_IDX["n"] += 1
            try:
                with open(os.path.join(API_DIR, f"raw_{RAW_IDX['n']:02d}.json"),
                          "w", encoding="utf-8") as f:
                    json.dump(data, f, ensure_ascii=False, indent=1)
            except Exception:
                pass
            dd = data.get("data") or {}
            orders = dd.get("orderList") or []
            if not orders:
                print(f"   [{rnd}] 接口返回空页 → 已到底")
                EMPTY_PAGE["hit"] = True
                break
            added = collect(data)
            nxt = dd.get("lastOrderTime")
            print(f"   [{rnd}] 本页 {len(orders)} 条，新增 {added}，累计 {len(ORDERS)}，"
                  f"下一页游标 {nxt}")
            if nxt in (None, cursor):
                print("      游标未推进，停止")
                break
            cursor = nxt
            pump(page, 1800)

        # ---- 兜底：接口翻页没走到空页时，再用滚动试 ----
        if not EMPTY_PAGE["hit"]:
            print(">>> 接口翻页未捕获空页，追加滚动兜底…")
            try:
                page.on("response", on_response)
            except Exception:
                pass
            stable = 0
            for _ in range(60):
                if EMPTY_PAGE["hit"]:
                    break
                before = len(ORDERS)
                do_scroll(page)
                pump(page, 3000)
                if len(ORDERS) > before:
                    stable = 0
                else:
                    stable += 1
                    if stable >= 6:
                        break
                    pump(page, 1500)
                if no_more_visible(page):
                    print("   页面出现「没有更多」提示，停止加载")
                    break

        if EMPTY_PAGE["hit"]:
            print(">>> 已确认到底（接口返回空页）")
        else:
            print("   !! 警告：未捕获到「空页」响应，列表可能仍有余量")

        filled = fill_categories()
        if filled:
            print(f">>> 分类补齐：{filled} 条（按 partnerid 实证映射）")
        rows = write_files()
        summary(rows)
        try:
            ctx.storage_state(path=STATE)
        except Exception:
            pass
        print(">>> 完成，按回车关闭浏览器…")
        try:
            input()
        except Exception:
            pass
        try:
            browser.close()
        except Exception:
            pass
        return rows


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    logf = open(LOG_PATH, "w", encoding="utf-8")
    sys.stdout = _Tee(sys.__stdout__, logf)
    sys.stderr = sys.stdout
    try:
        main()
    except KeyboardInterrupt:
        print("\n>>> 已中断（Ctrl+C），先把已抓到的订单落盘…")
        try:
            summary(write_files())
        except Exception:
            traceback.print_exc()
    except Exception:
        traceback.print_exc()
        try:
            summary(write_files())
        except Exception:
            pass
        try:
            input(">>> 出错了，按回车关闭…")
        except Exception:
            pass
    finally:
        try:
            print(f">>> 日志: {LOG_PATH}")
        except Exception:
            pass
        try:
            logf.close()
        except Exception:
            pass
